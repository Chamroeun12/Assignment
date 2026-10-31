import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = pptx.Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Tokens
    NAVY_DARK = RGBColor(15, 23, 42)       # #0F172A
    NAVY_CARD = RGBColor(30, 41, 59)       # #1E293B
    BLUE_ACCENT = RGBColor(2, 132, 199)    # #0284C7
    BLUE_LIGHT = RGBColor(224, 242, 254)   # #E0F2FE
    TEXT_DARK = RGBColor(30, 41, 59)       # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)   # #64748B
    TEXT_LIGHT = RGBColor(248, 250, 252)   # #F8FAFC
    CARD_BG = RGBColor(241, 245, 249)      # #F1F5F9
    BORDER_COLOR = RGBColor(203, 213, 225) # #CBD5E1
    GREEN_SUCCESS = RGBColor(22, 163, 74)  # #16A34A
    GREEN_LIGHT = RGBColor(220, 252, 231)  # #DCFCE7

    def add_header(slide, category, title, dark=False):
        # Category tracker tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category.upper()
        p_tag.font.name = "Segoe UI"
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = BLUE_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT if dark else NAVY_DARK

        # Underline bar
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(2.0), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = BLUE_ACCENT
        line.line.color.rgb = BLUE_ACCENT

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_DARK
    bg1.line.fill.background()

    # Title Card
    tcard = add_card(slide1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), bg_color=NAVY_CARD, border_color=BLUE_ACCENT)

    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(1.4), Inches(3.2), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_ACCENT
    badge.line.fill.background()
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "CLOUD COMPUTING CAPSTONE"
    p_b.font.name = "Segoe UI"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = RGBColor(255, 255, 255)
    p_b.alignment = PP_ALIGN.CENTER

    # Main Title
    t_box = slide1.shapes.add_textbox(Inches(1.3), Inches(2.0), Inches(10.5), Inches(2.2))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "Automated CI/CD, Infrastructure as Code &\nMulti-Service Container Orchestration"
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(32)
    p_t.font.bold = True
    p_t.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = tf_t.add_paragraph()
    p_sub.text = "Production Architecture on AWS EC2 with GitHub, Jenkins, Terraform, Docker Swarm & Nginx"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = BLUE_LIGHT

    # Metadata Grid
    meta_box = slide1.shapes.add_textbox(Inches(1.3), Inches(4.5), Inches(10.5), Inches(1.8))
    tf_m = meta_box.text_frame
    tf_m.word_wrap = True

    p_m1 = tf_m.paragraphs[0]
    p_m1.text = "Student / Engineer: Chamroeun Saran  |  Course: Master of Science in Cloud Computing"
    p_m1.font.name = "Segoe UI"
    p_m1.font.size = Pt(13)
    p_m1.font.bold = True
    p_m1.font.color.rgb = RGBColor(255, 255, 255)

    p_m2 = tf_m.add_paragraph()
    p_m2.text = "Repository: https://github.com/Chamroeun12/Assignment  |  AWS Region: Sydney (ap-southeast-2)"
    p_m2.font.name = "Segoe UI"
    p_m2.font.size = Pt(11)
    p_m2.font.color.rgb = TEXT_MUTED

    p_m3 = tf_m.add_paragraph()
    p_m3.text = "Verification: Jenkins Pipeline #4 Passed  |  Live Cluster IP: 16.178.178.44  |  October 2026"
    p_m3.font.name = "Segoe UI"
    p_m3.font.size = Pt(11)
    p_m3.font.bold = True
    p_m3.font.color.rgb = GREEN_SUCCESS

    # =========================================================================
    # SLIDE 2: Project Objectives & Requirements
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Capstone Scope & Objectives", "Engineering Goals & Rubric Fulfillment")

    cards_data2 = [
        ("1. GitOps & Version Control", [
            "Feature branching strategy (feature/eks-argocd-cicd)",
            "Pull Request workflow with code reviews and automated hooks",
            "Frequent, atomic commits with descriptive conventional messages",
            "Zero credentials or sensitive secrets committed to repository"
        ]),
        ("2. Infrastructure as Code (IaC)", [
            "100% reproducible AWS EC2 infrastructure using Terraform HCL",
            "Dedicated VPC, public subnets, internet gateway, and route tables",
            "Least-privilege Security Groups with internal cluster isolation (self=true)",
            "Dynamic Canonical Ubuntu 22.04 LTS AMI discovery"
        ]),
        ("3. Automated CI/CD Pipeline", [
            "Declarative Jenkinsfile pipeline triggered automatically on push",
            "Automated PEP8 linting (Flake8) and unit tests (Pytest 5/5, 89% cov)",
            "Container layer vulnerability scanning via Trivy CVE inspector",
            "Automated image building and publishing to Amazon ECR"
        ]),
        ("4. Resilient Container Orchestration", [
            "Docker Swarm multi-node cluster (1 Manager Leader + 2 Workers)",
            "Nginx reverse proxy load-balancing traffic across backend replicas",
            "Zero-downtime rolling update strategy (order: start-first)",
            "Extra Credit (+5): Automated Telegram alerts on build status"
        ])
    ]

    coords2 = [
        (Inches(0.8), Inches(1.8)),
        (Inches(6.8), Inches(1.8)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5))
    ]

    for (title, bullets), (left, top) in zip(cards_data2, coords2):
        add_card(slide2, left, top, Inches(5.7), Inches(2.4))
        tb = slide2.shapes.add_textbox(left + Inches(0.25), top + Inches(0.15), Inches(5.2), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT

        for bullet in bullets:
            p = tf.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_DARK
            p.space_after = Pt(2)

    # =========================================================================
    # SLIDE 3: System Architecture
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "System Architecture", "End-to-End Request Path & CI/CD Pipeline Topology")

    # Left Column: Key Components List
    add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.2), Inches(5.1))
    tb3 = slide3.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(4.7), Inches(4.7))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0

    p = tf3.paragraphs[0]
    p.text = "Architecture Components"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    arch_items = [
        ("GitHub Repository", "Authoritative source of truth for app code, Dockerfiles, and IaC."),
        ("Jenkins CI/CD", "Automates testing, CVE scanning, ECR pushing, and stack deployment."),
        ("Amazon ECR", "Private AES-256 encrypted registry hosting versioned container images."),
        ("AWS EC2 & VPC", "Dedicated 10.0.0.0/16 VPC with 1 Manager (t3.small) and 2 Workers (t3.micro)."),
        ("Nginx Load Balancer", "Reverse proxy on Port 80 distributing ingress across Flask replicas."),
        ("Docker Swarm Mesh", "Overlay network (Port 4789) routing requests via internal VIP (backend_app)."),
        ("Telegram Bot Alerts", "Instant push notifications reporting job status, tag, and build health.")
    ]

    for title, desc in arch_items:
        p_t = tf3.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK

        # Append desc in normal weight
        run = p_t.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_t.space_after = Pt(4)

    # Right Column: Architecture Diagram / Visual
    add_card(slide3, Inches(6.3), Inches(1.8), Inches(6.2), Inches(5.1))
    diag_path = "CICD_Pipeline_Commit_to_Production.jpg"
    if os.path.exists(diag_path):
        slide3.shapes.add_picture(diag_path, Inches(6.45), Inches(1.95), width=Inches(5.9))
    else:
        # Fallback diagram box
        dtb = slide3.shapes.add_textbox(Inches(6.5), Inches(2.0), Inches(5.8), Inches(4.5))
        dtf = dtb.text_frame
        dtf.word_wrap = True
        dp = dtf.paragraphs[0]
        dp.text = "CI/CD & Swarm Flow:\nDeveloper -> GitHub -> Jenkins -> ECR -> EC2 Swarm (Nginx:80 -> Flask:5000) -> Telegram"
        dp.font.name = "Courier New"
        dp.font.size = Pt(11)

    # =========================================================================
    # SLIDE 4: Infrastructure as Code with Terraform
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Infrastructure as Code", "Declarative AWS Provisioning via Terraform")

    # Left: Explanation & Highlights
    add_card(slide4, Inches(0.8), Inches(1.8), Inches(5.4), Inches(5.1))
    tb4 = slide4.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(4.9), Inches(4.7))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    tf4.margin_left = tf4.margin_top = tf4.margin_right = tf4.margin_bottom = 0

    p = tf4.paragraphs[0]
    p.text = "Terraform IaC Highlights"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    tf_points = [
        ("Modular Architecture", "Structured across main.tf, variables.tf, outputs.tf, and backend.tf for maintainability."),
        ("Dynamic AMI Resolution", "Uses data 'aws_ami' 'ubuntu' filter querying Canonical publisher (099720109477), eliminating hardcoded IDs."),
        ("Least-Privilege Security Group", "Restricts SSH (22) to admin CIDR, Web (80/443) to public, and Swarm (2377, 7946, 4789) strictly to self=true."),
        ("Multi-Node Compute Topology", "Provisions 1 Swarm Manager (t3.small) and 2 Workers (t3.micro) tagged with environment and role metadata."),
        ("State Management", "Configured S3 remote state storage with DynamoDB state locking to prevent concurrent mutation conflicts.")
    ]

    for title, desc in tf_points:
        p_t = tf4.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK
        run = p_t.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_t.space_after = Pt(5)

    # Right: Terraform Apply Screenshot
    add_card(slide4, Inches(6.5), Inches(1.8), Inches(6.0), Inches(5.1))
    tf_img = "screenshots/terraform_apply.png"
    if os.path.exists(tf_img):
        slide4.shapes.add_picture(tf_img, Inches(6.65), Inches(2.0), width=Inches(5.7))

    # =========================================================================
    # SLIDE 5: Docker Containerization Best Practices
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Containerization Standards", "Production Dockerfile & Security Hardening")

    # Left: Docker Practices
    add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.4), Inches(5.1))
    tb5 = slide5.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(4.9), Inches(4.7))
    tf5 = tb5.text_frame
    tf5.word_wrap = True
    tf5.margin_left = tf5.margin_top = tf5.margin_right = tf5.margin_bottom = 0

    p = tf5.paragraphs[0]
    p.text = "Containerization Standards"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    docker_points = [
        ("Multi-Stage Build Pattern", "Separates Python compilation dependencies (builder) from minimal production runtime (python:3.10-slim)."),
        ("Non-Root Security Context", "Runs as dedicated appuser (UID 1000) with chown -R appuser:appuser /app to mitigate container breakout risks."),
        ("Container Healthcheck", "Native Docker HEALTHCHECK validates endpoint http://127.0.0.1:5000/health every 30s with 3 retries."),
        ("Local Multi-Service Compose", "docker-compose.yml defines local developer environment with web and nginx services sharing a bridge network."),
        ("Minimal Image Footprint", "Utilizes .dockerignore to strip caches, git history, and secrets, reducing attack surface and pull latency.")
    ]

    for title, desc in docker_points:
        p_t = tf5.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK
        run = p_t.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_t.space_after = Pt(6)

    # Right: Local Container Preview
    add_card(slide5, Inches(6.5), Inches(1.8), Inches(6.0), Inches(5.1))
    app_img = "screenshots/app_container_healthcheck.png"
    if os.path.exists(app_img):
        slide5.shapes.add_picture(app_img, Inches(6.65), Inches(2.2), width=Inches(5.7))

    # =========================================================================
    # SLIDE 6: Automated Jenkins CI/CD Pipeline (Live Evidence)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "CI/CD Pipeline as Code", "Jenkins Declarative Pipeline — 100% Green Build #4")

    # Left: Pipeline Stages
    add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.2), Inches(5.1))
    tb6 = slide6.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(4.7), Inches(4.7))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    tf6.margin_left = tf6.margin_top = tf6.margin_right = tf6.margin_bottom = 0

    p = tf6.paragraphs[0]
    p.text = "Pipeline Stages & Metrics"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    stages = [
        ("1. Checkout SCM", "Clones latest commit from GitHub (Chamroeun12/Assignment.git) and dynamically creates version tag."),
        ("2. Build & Test", "Creates Python venv, runs Flake8 PEP 8 (0 errors) and Pytest (5/5 passed, 89% code coverage)."),
        ("3. Security & Quality Scan", "Scans container image layers for CVE vulnerabilities using Trivy inspector."),
        ("4. Authenticate & Push ECR", "Authenticates Docker client via aws ecr get-login-password and publishes versioned & latest tags."),
        ("5. Deploy Stack to Swarm", "Executes docker stack deploy --with-registry-auth to EC2 cluster with zero downtime.")
    ]

    for title, desc in stages:
        p_t = tf6.add_paragraph()
        p_t.text = f"✓ {title}: "
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = GREEN_SUCCESS
        run = p_t.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_t.space_after = Pt(4)

    # Right: Live Screenshot of Build #4
    add_card(slide6, Inches(6.3), Inches(1.8), Inches(6.2), Inches(5.1))
    jk_img = "screenshots/jenkins_build_success_live.png"
    if os.path.exists(jk_img):
        slide6.shapes.add_picture(jk_img, Inches(6.45), Inches(2.05), width=Inches(5.9))

    # =========================================================================
    # SLIDE 7: Docker Swarm Orchestration & Zero Downtime
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Container Orchestration", "Docker Swarm Multi-Node Topology & Rolling Updates")

    # Left: Explanation
    add_card(slide7, Inches(0.8), Inches(1.8), Inches(5.4), Inches(5.1))
    tb7 = slide7.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(4.9), Inches(4.7))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    tf7.margin_left = tf7.margin_top = tf7.margin_right = tf7.margin_bottom = 0

    p = tf7.paragraphs[0]
    p.text = "Swarm Cluster Architecture"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    swarm_points = [
        ("Cluster Formation", "1 Manager (Leader) managing Raft consensus (2377/tcp) and 2 Worker nodes in Ready & Active state."),
        ("Routing Mesh & Overlay Network", "VXLAN overlay (4789/udp) provides multi-host container networking with built-in virtual IP load balancing."),
        ("Zero-Downtime Rolling Update", "Configured order: start-first, parallelism: 1, delay: 10s. New container must pass healthcheck before terminating old task."),
        ("Credential Forwarding", "Uses docker stack deploy --with-registry-auth so worker nodes authenticate securely with Amazon ECR via the manager."),
        ("Fault Tolerance", "Automatic task rescheduling if any node or container fails, ensuring constant 2/2 replica availability.")
    ]

    for title, desc in swarm_points:
        p_t = tf7.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK
        run = p_t.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_t.space_after = Pt(5)

    # Right: Swarm nodes & stack ps screenshots
    add_card(slide7, Inches(6.5), Inches(1.8), Inches(6.0), Inches(5.1))
    sw_img = "screenshots/swarm_cluster_nodes.png"
    if os.path.exists(sw_img):
        slide7.shapes.add_picture(sw_img, Inches(6.65), Inches(2.2), width=Inches(5.7))

    # =========================================================================
    # SLIDE 8: Nginx Reverse Proxy & Load Balancing
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Network Ingress & Routing", "Nginx Reverse Proxy & Round-Robin Load Balancing")

    # Left: Ingress Configuration
    add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.4), Inches(5.1))
    tb8 = slide8.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(4.9), Inches(4.7))
    tf8 = tb8.text_frame
    tf8.word_wrap = True
    tf8.margin_left = tf8.margin_top = tf8.margin_right = tf8.margin_bottom = 0

    p = tf8.paragraphs[0]
    p.text = "Reverse Proxy Architecture"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    nginx_points = [
        ("Port 80 Ingress Termination", "Nginx binds to host Port 80, shielding backend application containers from direct public exposure."),
        ("Round-Robin Traffic Balancing", "Routes requests to internal VIP upstream backend_app (web:5000) across replicated Flask tasks."),
        ("GET / (Web UI)", "Renders interactive cloud dashboard showing hostname, system stats, and runtime metrics."),
        ("GET /health (Application Health)", "HTTP 200 OK returning {\"status\": \"healthy\", \"version\": \"1.0.0\"}."),
        ("GET /api/info (Backend Metadata)", "HTTP 200 OK displaying container hostname, cluster type (docker-swarm), and runtime."),
        ("GET /nginx-health", "Dedicated proxy health probe confirming reverse proxy routing integrity.")
    ]

    for title, desc in nginx_points:
        p_t = tf8.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK
        run = p_t.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_t.space_after = Pt(4)

    # Right: Live Curl verification screenshot
    add_card(slide8, Inches(6.5), Inches(1.8), Inches(6.0), Inches(5.1))
    curl_img = "screenshots/nginx_load_balancer_curl.png"
    if os.path.exists(curl_img):
        slide8.shapes.add_picture(curl_img, Inches(6.65), Inches(2.2), width=Inches(5.7))

    # =========================================================================
    # SLIDE 9: Extra Credit — Telegram Notifications
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Operational Observability (+5 Extra Credit)", "Real-Time CI/CD Notifications via Telegram Bot API")

    # Left: Telegram Implementation Details
    add_card(slide9, Inches(0.8), Inches(1.8), Inches(5.4), Inches(5.1))
    tb9 = slide9.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(4.9), Inches(4.7))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    tf9.margin_left = tf9.margin_top = tf9.margin_right = tf9.margin_bottom = 0

    p = tf9.paragraphs[0]
    p.text = "Telegram Alerting System"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    tg_points = [
        ("Automated Post-Action Callbacks", "Integrated directly into Jenkinsfile post { success { ... } failure { ... } } blocks."),
        ("Distinct Success & Failure Alerts", "Sends formatted markdown messages reporting Job name, Build #, Git SHA, and Target cluster."),
        ("Secure Credential Injection", "Bot token and Chat ID are fetched dynamically via Jenkins withCredentials, preventing credential leaks."),
        ("Non-Blocking Resilient Design", "Wrapped in try/catch with fallback logging so builds succeed even if notifications are skipped."),
        ("Mobile & Desktop Instant Delivery", "DevOps engineers receive push notifications seconds after pipeline completion on mobile devices.")
    ]

    for title, desc in tg_points:
        p_t = tf9.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK
        run = p_t.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p_t.space_after = Pt(5)

    # Right: Telegram notification screenshot
    add_card(slide9, Inches(6.5), Inches(1.8), Inches(6.0), Inches(5.1))
    tg_img = "screenshots/telegram_notification.png"
    if os.path.exists(tg_img):
        slide9.shapes.add_picture(tg_img, Inches(6.65), Inches(2.2), width=Inches(5.7))

    # =========================================================================
    # SLIDE 10: Engineering Challenges & Implemented Solutions
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Engineering Rigor", "Key Technical Challenges & Production Solutions")

    challenges = [
        ("1. Swarm Worker ECR Pull Authentication",
         "Workers pull images directly from Amazon ECR but lacked AWS credentials stored on Manager.",
         "Added --with-registry-auth flag to docker stack deploy, which forwards the Manager's ECR auth token securely through Swarm Raft control plane."),
        ("2. Zero Downtime During Redeployments",
         "Default Docker redeployment terminates existing tasks before initializing new ones (order: stop-first).",
         "Configured order: start-first and parallelism: 1 with container healthcheck to guarantee new replicas are healthy before killing old tasks."),
        ("3. Dynamic AMI Portability Across AWS Regions",
         "Hardcoded AMI IDs cause failure when running Terraform in different regions or after Ubuntu updates.",
         "Replaced static IDs with data 'aws_ami' 'ubuntu' querying Canonical publisher ID (099720109477) for latest 22.04 LTS build."),
        ("4. Restricting Docker Swarm Cluster Port Exposure",
         "Exposing Swarm management (2377) and VXLAN (4789) to 0.0.0.0/0 introduces major security risks.",
         "Configured Terraform security group ingress rules with self=true, isolating cluster communication strictly to nodes inside the security group.")
    ]

    coords10 = [
        (Inches(0.8), Inches(1.8)),
        (Inches(6.8), Inches(1.8)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5))
    ]

    for (title, problem, solution), (left, top) in zip(challenges, coords10):
        add_card(slide10, left, top, Inches(5.7), Inches(2.4))
        tb = slide10.shapes.add_textbox(left + Inches(0.25), top + Inches(0.15), Inches(5.2), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT

        p1 = tf.add_paragraph()
        p1.text = "Challenge: "
        p1.font.bold = True
        p1.font.size = Pt(9.5)
        p1.font.color.rgb = RGBColor(220, 38, 38)
        run1 = p1.add_run()
        run1.text = problem
        run1.font.bold = False
        run1.font.color.rgb = TEXT_DARK
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = "Solution: "
        p2.font.bold = True
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = GREEN_SUCCESS
        run2 = p2.add_run()
        run2.text = solution
        run2.font.bold = False
        run2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 11: Live Demonstration Guide for Examiner
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "Demonstration & Verification", "Step-by-Step Live Demonstration for Lecturer")

    demo_steps = [
        ("Step 1: Live Web Application", [
            "Open browser: http://16.178.178.44",
            "Demonstrates Nginx reverse proxy serving the full interactive dashboard on Port 80.",
            "Verify real-time container hostname and system statistics."
        ]),
        ("Step 2: HTTP Endpoint Curl Verification", [
            "curl -s http://16.178.178.44/health  -> Returns {\"status\": \"healthy\", \"version\": \"1.0.0\"}",
            "curl -s http://16.178.178.44/api/info -> Demonstrates round-robin hostname switching across replicas",
            "curl -s http://16.178.178.44/nginx-health -> Confirms load balancer routing integrity"
        ]),
        ("Step 3: Jenkins CI/CD Automation", [
            "Open Jenkins: http://localhost:8080/job/Assignment-CI-CD/",
            "Show Build #4 Pipeline Stage View with all 5 stages green.",
            "Inspect console output confirming Pytest (5/5 passed, 89% cov) and Flake8 (0 errors)."
        ]),
        ("Step 4: Docker Swarm Cluster via SSH", [
            "ssh -i chamroeunsaran.pem ubuntu@16.178.178.44",
            "Run 'sudo docker service ls' to demonstrate 2/2 backend and 1/1 nginx replicas.",
            "Run 'sudo docker stack ps cicd-assignment' showing distributed tasks across cluster nodes."
        ])
    ]

    coords11 = [
        (Inches(0.8), Inches(1.8)),
        (Inches(6.8), Inches(1.8)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5))
    ]

    for (title, bullets), (left, top) in zip(demo_steps, coords11):
        add_card(slide11, left, top, Inches(5.7), Inches(2.4))
        tb = slide11.shapes.add_textbox(left + Inches(0.25), top + Inches(0.15), Inches(5.2), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT

        for b in bullets:
            p = tf.add_paragraph()
            p.text = f"•  {b}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(9.5)
            p.font.color.rgb = TEXT_DARK
            p.space_after = Pt(2)

    # =========================================================================
    # SLIDE 12: Conclusion & Deliverables Summary (Dark Theme)
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    bg12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = NAVY_DARK
    bg12.line.fill.background()

    card12 = add_card(slide12, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), bg_color=NAVY_CARD, border_color=BLUE_ACCENT)

    # Title in card
    tb12 = slide12.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(10.5), Inches(5.0))
    tf12 = tb12.text_frame
    tf12.word_wrap = True
    tf12.margin_left = tf12.margin_top = tf12.margin_right = tf12.margin_bottom = 0

    p_c0 = tf12.paragraphs[0]
    p_c0.text = "PROJECT DELIVERABLES & CONCLUSION"
    p_c0.font.name = "Segoe UI"
    p_c0.font.size = Pt(12)
    p_c0.font.bold = True
    p_c0.font.color.rgb = BLUE_ACCENT

    p_c1 = tf12.add_paragraph()
    p_c1.text = "All Capstone Requirements 100% Verified & Fulfilled"
    p_c1.font.name = "Segoe UI"
    p_c1.font.size = Pt(24)
    p_c1.font.bold = True
    p_c1.font.color.rgb = RGBColor(255, 255, 255)
    p_c1.space_after = Pt(12)

    summary_bullets = [
        ("Git Repository", "https://github.com/Chamroeun12/Assignment (Clean commit history & feature branches)"),
        ("CI/CD Pipeline", "Jenkins declarative pipeline with 5 automated stages (All Green in Build #4)"),
        ("Infrastructure as Code", "Complete Terraform modules with dynamic AMI, VPC, SGs, and EC2 nodes"),
        ("Container Orchestration", "Multi-node Docker Swarm with zero-downtime rolling updates & Nginx proxy"),
        ("Official Report Deliverables", "REPORT.pdf (4 pages, embedded evidence) & REPORT.docx formatted reports"),
        ("Extra Credit (+5 Points)", "Automated Telegram Bot push notification alerts configured and verified")
    ]

    for label, val in summary_bullets:
        p_sb = tf12.add_paragraph()
        p_sb.text = f"✓ {label}: "
        p_sb.font.name = "Segoe UI"
        p_sb.font.size = Pt(11)
        p_sb.font.bold = True
        p_sb.font.color.rgb = GREEN_SUCCESS
        run = p_sb.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = TEXT_LIGHT
        p_sb.space_after = Pt(4)

    p_end = tf12.add_paragraph()
    p_end.text = "\nThank You!  —  Questions & Discussion"
    p_end.font.name = "Segoe UI"
    p_end.font.size = Pt(16)
    p_end.font.bold = True
    p_end.font.color.rgb = BLUE_ACCENT
    p_end.alignment = PP_ALIGN.CENTER

    # Save presentation
    output_path = "Assignment.pptx"
    prs.save(output_path)
    print(f"Successfully generated presentation: {output_path}")

if __name__ == "__main__":
    create_presentation()
